import csv
import io
from calendar import monthrange
from dataclasses import dataclass
from decimal import Decimal
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

from .domain import Order, OrderItem, OrderStatus, Product, Role, User


@dataclass(slots=True)
class ImportedDataset:
    products: list[Product]
    users: list[User]
    orders: list[Order]


def load_dataset(import_dir: Path) -> ImportedDataset:
    products = load_products(import_dir / "Tovar.xlsx")
    users = load_users(import_dir / "user_import.xlsx")
    # строим словарь «полное имя → логин» для сопоставления заказов
    name_to_login: dict[str, str] = {u.full_name: u.login for u in users}
    orders = load_orders(import_dir / "Заказ_import.xlsx", products, name_to_login)
    return ImportedDataset(products=products, users=users, orders=orders)


def load_products(path: Path) -> list[Product]:
    """загрузка товаров из файла excel по пути."""
    with open(path, "rb") as f:
        return generic_load_products(f, path.name)


def generic_load_products(file_stream, filename: str) -> list[Product]:
    """универсальный загрузчик товаров из потока данных (csv или xlsx)."""
    if filename.lower().endswith((".xlsx", ".xls")):
        wb = load_workbook(io.BytesIO(file_stream.read()), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        data_rows = rows[1:]
    else:
        # считаем что это csv
        content = file_stream.read().decode("utf-8")
        reader = csv.reader(io.StringIO(content), delimiter=";")
        rows = list(reader)
        data_rows = rows[1:]

    products: list[Product] = []
    for row in data_rows:
        if not row or len(row) < 4 or not row[0]:
            continue
        try:
            # вспомогательные функции для безопасного парсинга
            def to_int(val):
                if val is None: return 0
                if isinstance(val, (int, float)): return int(val)
                s = str(val).strip()
                if not s: return 0
                try:
                    return int(float(s.replace(",", ".")))
                except ValueError:
                    return 0

            def to_decimal(val):
                if val is None: return Decimal("0")
                if isinstance(val, (int, float, Decimal)): return Decimal(str(val))
                s = str(val).strip().replace(",", ".").replace(" ", "")
                if not s: return Decimal("0")
                try:
                    return Decimal(s)
                except Exception:
                    return Decimal("0")

            products.append(
                Product(
                    article=str(row[0]).strip(),
                    name=str(row[1]).strip() if len(row) > 1 else "без названия",
                    unit=str(row[2]).strip() if len(row) > 2 else "шт.",
                    price=to_decimal(row[3]) if len(row) > 3 else Decimal("0"),
                    supplier=str(row[4]).strip() if len(row) > 4 else "неизвестно",
                    manufacturer=str(row[5]).strip() if len(row) > 5 else "неизвестно",
                    category=str(row[6]).strip() if len(row) > 6 else "общая",
                    discount_percent=to_int(row[7]) if len(row) > 7 else 0,
                    stock_count=to_int(row[8]) if len(row) > 8 else 0,
                    description=str(row[9]).strip() if len(row) > 9 and row[9] else "",
                    image_path=str(row[10]).strip() if len(row) > 10 and row[10] else None,
                    max_discount_percent=15, # в файле нет этого поля, ставим дефолт
                )
            )
        except Exception:
            # всё равно пропускаем битые строки, но теперь их будет меньше
            continue
    return products


def load_users(path: Path) -> list[User]:
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    users: list[User] = []
    for row in rows[1:]:
        if not row or not row[1]:
            continue
        role = normalize_role(str(row[0] or "").strip())
        users.append(
            User(
                login=str(row[2]).strip(),
                full_name=str(row[1]).strip(),
                role=role,
                password_hash=str(row[3]).strip(),
            )
        )
    return users


def load_orders(
    path: Path,
    products: list[Product],
    name_to_login: dict[str, str],
) -> list[Order]:
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    catalog = {product.article: product for product in products}
    orders: list[Order] = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        full_name = str(row[4] or "").strip()
        login = name_to_login.get(full_name, "")
        lines = parse_order_lines(str(row[1] or ""))
        items = [
            OrderItem(article=article, quantity=quantity, unit_price=catalog[article].discounted_price)
            for article, quantity in lines
            if article in catalog
        ]
        orders.append(
            Order(
                number=int(row[0]),
                user_login=login,
                created_at=normalize_date(row[2]),
                delivery_date=normalize_date(row[3]),
                pickup_code=int(row[5]),
                status=normalize_status(row[6]),
                items=items,
            )
        )
    return orders


def parse_order_lines(text: str) -> list[tuple[str, int]]:
    parts = [part.strip() for part in text.split(",") if part.strip()]
    result: list[tuple[str, int]] = []
    for index in range(0, len(parts), 2):
        article = parts[index]
        quantity = int(parts[index + 1]) if index + 1 < len(parts) else 1
        result.append((article, quantity))
    return result


def normalize_role(value: str) -> Role:
    value = value.casefold()
    if "админ" in value:
        return Role.ADMIN
    if "менедж" in value:
        return Role.MANAGER
    return Role.CLIENT


def normalize_status(value: str) -> OrderStatus:
    value = value.strip().casefold()
    if value.startswith("нов"):
        return OrderStatus.NEW
    if value.startswith("зав"):
        return OrderStatus.DONE
    return OrderStatus.PROCESSING


def normalize_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    if hasattr(value, "date"):
        return value.date()
    text = str(value).strip()
    if not text:
        return date.today()
    for parser in (_parse_iso_date, _parse_ru_date):
        parsed = parser(text)
        if parsed is not None:
            return parsed
    return date.today()


def _parse_iso_date(text: str) -> date | None:
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _parse_ru_date(text: str) -> date | None:
    parts = text.split(".")
    if len(parts) != 3:
        return None
    try:
        day = int(parts[0])
        month = int(parts[1])
        year = int(parts[2])
    except ValueError:
        return None
    month = max(1, min(month, 12))
    last_day = monthrange(year, month)[1]
    day = max(1, min(day, last_day))
    return date(year, month, day)


# lookup_login удалена: логины в user_import.xlsx — это email-адреса,
# а не транслитерация ФИО. сопоставление теперь идёт через name_to_login.
